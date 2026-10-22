import io
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from .models import Bureau, Centre, MoiAussi, Module, ProfilAgent, Remarque, VersionA3H

User = get_user_model()
DOSSIER = tempfile.mkdtemp()
PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
       b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe"
       b"\x02\xfe\xa7\x9a\xa0\xa0\x00\x00\x00\x00IEND\xaeB`\x82")


@override_settings(REMARQUES_CAPTURES_DIR=DOSSIER)
class RemarquesTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(DOSSIER, ignore_errors=True)

    def setUp(self):
        self.centre = Centre.objects.create(nom="Centre Tunis")
        self.autre_centre = Centre.objects.create(nom="Centre Sfax")
        self.bureau = Bureau.objects.create(centre=self.centre, nom="Bureau 1")
        self.bureau2 = Bureau.objects.create(centre=self.autre_centre, nom="Bureau 2")
        self.module = Module.objects.get(nom="Calcul réel")
        self.a = User.objects.create_user("agent_a", password="x")
        self.b = User.objects.create_user("agent_b", password="x")
        self.admin = User.objects.create_user("chef", password="x", is_staff=True, is_superuser=True)
        ProfilAgent.objects.create(user=self.a, nom_agent="Agent A", bureau=self.bureau)
        ProfilAgent.objects.create(user=self.b, nom_agent="Agent B", bureau=self.bureau2)

    def _donnees(self, **extra):
        d = {"nom_agent": "Agent A", "module": self.module.pk, "type_remarque": "calcul",
             "description": "Le total est faux sur le code 12040"}
        d.update(extra)
        return d

    def test_modules_par_defaut(self):
        self.assertGreaterEqual(Module.objects.count(), 15)

    def test_pages_autonomes_sans_gabarit_du_site(self):
        self.client.force_login(self.a)
        for nom in ("accueil", "nouvelle", "mes", "idees", "profil", "nouveautes"):
            page = self.client.get(reverse("remarques:" + nom)).content.decode()
            self.assertIn("Remarques des agents", page, nom)   # en-tête propre au module
            self.assertNotIn("Accueil du site", page, nom)      # ne dépend plus du base.html du site

    def test_accueil_compteurs_et_nouveautes(self):
        v = VersionA3H.objects.create(numero="2.5", notes="Correction des arrondis")
        Remarque.objects.create(
            auteur=self.a, nom_agent="Agent A", bureau=self.bureau, module=self.module,
            type_remarque="arrondi", description="arrondi faux", statut="corrigee",
            version_correction=v, resume_public="Arrondi au dinar corrigé sur le code 12040",
            accepte_credit=True)
        Remarque.objects.create(
            auteur=self.b, nom_agent="Agent B", bureau=self.bureau2, module=self.module,
            type_remarque="calcul", description="autre", statut="corrigee",
            version_correction=v, resume_public="Total mensuel corrigé")
        self.client.force_login(self.a)
        page = self.client.get(reverse("remarques:accueil")).content.decode()
        self.assertIn("Arrondi au dinar corrigé", page)
        self.assertIn("Agent A, Bureau 1", page)       # crédité : accord donné
        self.assertIn("Bureau Bureau 2", page)         # non crédité : bureau seulement
        self.assertNotIn("Agent B", page.split("Ce qui a changé")[1])
        self.assertIn("Merci Agent A", page)
        nv = self.client.get(reverse("remarques:nouveautes")).content.decode()
        self.assertIn("Version 2.5", nv)
        self.assertIn("Correction des arrondis", nv)

    def test_tuile_preselectionne_le_type(self):
        self.client.force_login(self.a)
        page = self.client.get(reverse("remarques:nouvelle") + "?type=arrondi").content.decode()
        self.assertRegex(page, r'value="arrondi"\s+selected')

    def test_badge_nouvelle_reponse_puis_lu(self):
        r = Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                    type_remarque="calcul", description="calcul faux")
        self.client.force_login(self.admin)
        champs = {"statut": "etudiee", "reponse": "Je regarde", "module": self.module.pk,
                  "type_remarque": "calcul", "description": "calcul faux", "impact": "",
                  "resume_public": "", "note_interne": ""}
        self.client.post(reverse("admin:remarques_remarque_change", args=[r.pk]), champs)
        r.refresh_from_db()
        self.assertTrue(r.non_lue)
        self.client.force_login(self.a)
        self.assertIn("Nouvelle réponse", self.client.get(reverse("remarques:mes")).content.decode())
        self.assertIn("1 nouvelle réponse", self.client.get(reverse("remarques:accueil")).content.decode())
        self.client.get(reverse("remarques:detail", args=[r.pk]))   # l'agent lit
        r.refresh_from_db()
        self.assertFalse(r.non_lue)
        self.assertNotIn("Nouvelle réponse", self.client.get(reverse("remarques:mes")).content.decode())

    def test_frise_d_avancement(self):
        from .views import _frise
        self.assertEqual([e["etat"] for e in _frise("nouvelle")], ["courant", "a_venir", "a_venir", "a_venir"])
        self.assertEqual([e["etat"] for e in _frise("planifiee")], ["fait", "fait", "courant", "a_venir"])
        self.assertEqual([e["etat"] for e in _frise("corrigee")], ["fait", "fait", "fait", "fait"])

    def test_connexion_et_deconnexion(self):
        r = self.client.get("/")
        self.assertRedirects(r, "/login/?next=/", fetch_redirect_response=False)
        page = self.client.get("/login/").content.decode()
        self.assertIn("Se connecter", page)
        self.assertNotIn("Nouvelle remarque", page)          # pas de menu sur la page de connexion
        mauvais = self.client.post("/login/", {"username": "agent_a", "password": "faux"})
        self.assertContains(mauvais, "incorrect")
        bon = self.client.post("/login/", {"username": "agent_a", "password": "x"})
        self.assertRedirects(bon, "/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/").status_code, 200)
        self.client.post("/logout/")
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_acces_anonyme_refuse(self):
        for nom in ("accueil", "nouvelle", "mes", "idees", "profil"):
            r = self.client.get(reverse("remarques:" + nom))
            self.assertEqual(r.status_code, 302, nom)

    def test_premiere_visite_demande_le_profil(self):
        u = User.objects.create_user("nouveau", password="x")
        self.client.force_login(u)
        r = self.client.get(reverse("remarques:nouvelle"))
        self.assertRedirects(r, reverse("remarques:profil") + "?next=" + reverse("remarques:nouvelle"),
                             fetch_redirect_response=False)

    def test_profil_enregistre_et_bureau_coherent(self):
        u = User.objects.create_user("nouveau", password="x")
        self.client.force_login(u)
        ko = self.client.post(reverse("remarques:profil"),
                              {"nom_agent": "N", "centre": self.centre.pk, "bureau": self.bureau2.pk})
        self.assertEqual(ko.status_code, 200)
        ok = self.client.post(reverse("remarques:profil"),
                              {"nom_agent": "N", "centre": self.centre.pk, "bureau": self.bureau.pk})
        self.assertEqual(ok.status_code, 302)
        self.assertEqual(u.profil_remarques.bureau, self.bureau)

    def test_formulaire_leger_trois_champs(self):
        self.client.force_login(self.a)
        r = self.client.post(reverse("remarques:nouvelle"), self._donnees())
        rem = Remarque.objects.get()
        self.assertRedirects(r, reverse("remarques:detail", args=[rem.pk]))
        self.assertEqual(rem.bureau, self.bureau)
        self.assertEqual(rem.auteur, self.a)
        self.assertEqual(rem.statut, "nouvelle")

    def test_champs_obligatoires(self):
        self.client.force_login(self.a)
        r = self.client.post(reverse("remarques:nouvelle"), self._donnees(description="  "))
        self.assertEqual(r.status_code, 200)
        r = self.client.post(reverse("remarques:nouvelle"), self._donnees(description="ok"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Remarque.objects.count(), 0)

    def test_texte_arabe_et_details(self):
        self.client.force_login(self.a)
        self.client.post(reverse("remarques:nouvelle"), self._donnees(
            description="الخصم من المورد غير صحيح", resultat_obtenu="960000",
            resultat_attendu="959999", source_attendu="manuel", code_impot="11041"))
        rem = Remarque.objects.get()
        page = self.client.get(reverse("remarques:detail", args=[rem.pk])).content.decode()
        self.assertIn("الخصم من المورد", page)
        self.assertIn("959999", page)

    def test_mes_remarques_prives(self):
        r = Remarque.objects.create(auteur=self.a, nom_agent="Agent A", bureau=self.bureau,
                                    module=self.module, type_remarque="calcul", description="secret calcul")
        self.client.force_login(self.b)
        self.assertNotIn("secret calcul", self.client.get(reverse("remarques:mes")).content.decode())
        self.assertEqual(self.client.get(reverse("remarques:detail", args=[r.pk])).status_code, 404)
        self.client.force_login(self.a)
        self.assertIn("secret calcul", self.client.get(reverse("remarques:mes")).content.decode())

    def test_reponse_visible_par_l_agent(self):
        v = VersionA3H.objects.create(numero="2.4")
        r = Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                    type_remarque="arrondi", description="arrondi faux",
                                    statut="corrigee", reponse="Corrigé, merci", version_correction=v)
        self.client.force_login(self.a)
        page = self.client.get(reverse("remarques:detail", args=[r.pk])).content.decode()
        self.assertIn("Corrigé, merci", page)
        self.assertIn("2.4", page)

    def test_idees_et_moi_aussi(self):
        idee = Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                       type_remarque="idee", description="Bouton export Excel")
        doublon = Remarque.objects.create(auteur=self.b, nom_agent="B", bureau=self.bureau2, module=self.module,
                                          type_remarque="idee", description="Exporter en excel", doublon_de=idee)
        c = User.objects.create_user("c", password="x")
        self.client.force_login(c)
        url = reverse("remarques:moi_aussi", args=[idee.pk])
        self.client.post(url)
        self.assertEqual(MoiAussi.objects.count(), 1)
        page = self.client.get(reverse("remarques:idees")).content.decode()
        self.assertEqual(page.count("Bouton export Excel"), 1)
        self.assertNotIn("Exporter en excel", page)  # doublon fusionné
        self.client.post(url)  # deuxième clic = retrait
        self.assertEqual(MoiAussi.objects.count(), 0)
        # l'auteur ne peut pas voter pour sa propre idée
        self.client.force_login(self.a)
        self.client.post(url)
        self.assertEqual(MoiAussi.objects.count(), 0)
        # un calcul n'est pas votable
        calc = Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                       type_remarque="calcul", description="x calcul")
        self.client.force_login(c)
        self.assertEqual(self.client.post(reverse("remarques:moi_aussi", args=[calc.pk])).status_code, 404)

    def test_tableau_reserve_aux_administrateurs(self):
        self.client.force_login(self.a)
        self.assertEqual(self.client.get(reverse("remarques:tableau")).status_code, 302)
        Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                type_remarque="calcul", description="calcul faux")
        self.client.force_login(self.admin)
        page = self.client.get(reverse("remarques:tableau"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Centre Tunis")

    def test_capture_protegee(self):
        self.client.force_login(self.a)
        f = SimpleUploadedFile("ecran.png", PNG, content_type="image/png")
        self.client.post(reverse("remarques:nouvelle"), self._donnees(type_remarque="formulaire", capture=f))
        rem = Remarque.objects.get()
        self.assertTrue(rem.capture)
        url = reverse("remarques:capture", args=[rem.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.force_login(self.b)
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_capture_format_refuse(self):
        self.client.force_login(self.a)
        f = SimpleUploadedFile("x.html", b"<script>alert(1)</script>", content_type="text/html")
        self.client.post(reverse("remarques:nouvelle"), self._donnees(capture=f))
        self.assertEqual(Remarque.objects.count(), 0)

    def test_admin_liste_filtres_et_export_excel(self):
        for i in range(3):
            Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                    type_remarque="calcul", description="calcul %d" % i)
        self.client.force_login(self.admin)
        liste = reverse("admin:remarques_remarque_changelist")
        self.assertEqual(self.client.get(liste).status_code, 200)
        self.assertEqual(self.client.get(liste + "?bureau__centre__id__exact=%d&statut__exact=nouvelle" % self.centre.pk).status_code, 200)
        ids = list(Remarque.objects.values_list("pk", flat=True))
        r = self.client.post(liste, {"action": "exporter_excel", "_selected_action": ids})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 4)
        r2 = self.client.get(reverse("admin:remarques_remarque_change", args=[ids[0]]))
        self.assertEqual(r2.status_code, 200)

    def test_admin_action_cas_de_test(self):
        r = Remarque.objects.create(auteur=self.a, nom_agent="A", bureau=self.bureau, module=self.module,
                                    type_remarque="arrondi", description="arrondi")
        self.client.force_login(self.admin)
        self.client.post(reverse("admin:remarques_remarque_changelist"),
                         {"action": "marquer_cas_test", "_selected_action": [r.pk]})
        r.refresh_from_db()
        self.assertTrue(r.cas_de_test)


class InitAdminTests(TestCase):
    def test_cree_le_compte_une_seule_fois_sans_toucher_au_mot_de_passe(self):
        from io import StringIO
        from unittest import mock
        from django.core.management import call_command
        env = {"A3H_ADMIN_USER": "chef2", "A3H_ADMIN_PASSWORD": "mdp-un"}
        with mock.patch.dict("os.environ", env):
            call_command("init_admin", stdout=StringIO())
            u = User.objects.get(username="chef2")
            self.assertTrue(u.is_superuser and u.check_password("mdp-un"))
            with mock.patch.dict("os.environ", {"A3H_ADMIN_PASSWORD": "autre"}):
                call_command("init_admin", stdout=StringIO())
            u.refresh_from_db()
            self.assertTrue(u.check_password("mdp-un"))   # jamais écrasé
        call_command("init_admin", stdout=StringIO())      # sans variables : ne fait rien
